"""Render a Markdown doc in this repo to a shareable .docx (figures embedded).

Usage:
    python scripts/make_docx.py                                                 # spectral_graph/INSIGHTS.md
    python scripts/make_docx.py experiments/boosting_trees/EXPERIMENT_INDEX.md   # any other doc
    python scripts/make_docx.py experiments/spectral_graph/INSIGHTS.md -o Out.docx

Handles the subset of Markdown these docs use: headings, wrapped paragraphs, pipe
tables, images with italic captions, horizontal rules, links, and inline
**bold** / *italic* / `code`.
"""
import argparse
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
INK = RGBColor(0x1A, 0x1A, 0x1A)
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)

HEADING_PT = {1: 21, 2: 13, 3: 11.5}

INLINE = re.compile(
    r"(\*\*.+?\*\*|(?<!\*)\*[^*]+?\*(?!\*)|`[^`]+?`|\[[^\]]+\]\([^)]+\))"
)
LINK = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")


# ----------------------------------------------------------------- helpers
def add_runs(par, text):
    """Write text into a paragraph, honouring **bold**, *italic*, `code`, links."""
    for piece in INLINE.split(text):
        if not piece:
            continue

        m = LINK.match(piece)
        if m:
            label, url = m.group(1), m.group(2)
            # A path is more useful to a reader than a redundant label.
            shown = url if "/" in url else label.strip("`")
            r = par.add_run(shown)
            r.font.name = MONO_FONT
            r.font.size = Pt(9)
            r.font.color.rgb = MUTED
            continue

        if piece.startswith("**") and piece.endswith("**"):
            par.add_run(piece[2:-2]).bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            r = par.add_run(piece[1:-1])
            r.font.name = MONO_FONT
            r.font.size = Pt(9.5)
        elif piece.startswith("*") and piece.endswith("*"):
            par.add_run(piece[1:-1]).italic = True
        else:
            par.add_run(piece)


def horizontal_rule(doc):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(10)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:color"), "D0D0D0")
    pbdr.append(bottom)
    par._p.get_or_add_pPr().append(pbdr)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_fill)
    tcPr.append(el)


def build_table(doc, rows, content_width):
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # size columns by their longest cell, with a floor so narrow columns stay legible
    spans = [max(len(r[c]) for r in rows) for c in range(ncols)]
    spans = [max(s, 6) for s in spans]
    total = sum(spans)
    widths = [int(content_width * s / total) for s in spans]

    for r, row in enumerate(rows):
        trPr = table.rows[r]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))  # no row torn across a page
        if r == 0:
            trPr.append(OxmlElement("w:tblHeader"))  # repeat header on each page
        for c, text in enumerate(row):
            cell = table.cell(r, c)
            cell.width = widths[c]
            par = cell.paragraphs[0]
            par.paragraph_format.space_before = Pt(3)
            par.paragraph_format.space_after = Pt(3)
            add_runs(par, text)
            for run in par.runs:
                run.font.size = Pt(9.5)
                if r == 0:
                    run.bold = True
            if r == 0:
                shade(cell, "EDF2F7")
    return table


# ----------------------------------------------------------------- document
def convert(src, out):
    src_dir = os.path.dirname(src)
    with open(src, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()

    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.85)
    section.top_margin = section.bottom_margin = Inches(0.8)
    content_width = section.page_width - section.left_margin - section.right_margin

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.12

    i = 0
    last_was_image = False
    first_heading_done = False

    while i < len(lines):
        line = lines[i].rstrip()

        if not line.strip():
            i += 1
            continue

        if re.fullmatch(r"-{3,}", line.strip()):
            horizontal_rule(doc)
            i += 1
            continue

        # --- headings
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(2)
                run = par.add_run(text)
                run.font.size = Pt(HEADING_PT[1])
                run.bold = True
                run.font.color.rgb = ACCENT
            else:
                par = doc.add_heading(level=level)
                par.paragraph_format.space_before = Pt(
                    (16 if level == 2 else 11) if first_heading_done else 8
                )
                par.paragraph_format.space_after = Pt(4)
                par.paragraph_format.keep_with_next = True
                for piece in INLINE.split(text):
                    if not piece:
                        continue
                    if piece.startswith("*") and piece.endswith("*") and not piece.startswith("**"):
                        run = par.add_run(piece[1:-1])
                        run.italic = True
                        run.font.color.rgb = MUTED
                    else:
                        run = par.add_run(piece)
                        run.bold = True
                        run.font.color.rgb = ACCENT
                    run.font.name = BODY_FONT
                    run.font.size = Pt(HEADING_PT[level])
                first_heading_done = True
            last_was_image = False
            i += 1
            continue

        # --- image
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)\s*$", line)
        if m:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(6)
            par.paragraph_format.space_after = Pt(2)
            par.paragraph_format.keep_with_next = True  # caption must not orphan
            par.add_run().add_picture(os.path.join(src_dir, m.group(2)), width=content_width)
            last_was_image = True
            i += 1
            continue

        # --- table block
        if line.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                build_table(doc, rows, content_width)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
            last_was_image = False
            continue

        # --- wrapped paragraph
        block = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,3}\s|!\[|\||→|-{3,}$)", lines[i].lstrip()
        ):
            block.append(lines[i].strip())
            i += 1
        text = " ".join(block)

        # figure caption, or the subtitle under the title
        if re.fullmatch(r"\*[^*].*\*", text):
            par = doc.add_paragraph()
            par.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if last_was_image else WD_ALIGN_PARAGRAPH.LEFT
            )
            par.paragraph_format.space_after = Pt(14 if last_was_image else 12)
            run = par.add_run(text[1:-1])
            run.italic = True
            run.font.size = Pt(9 if last_was_image else 10.5)
            run.font.color.rgb = MUTED
            last_was_image = False
            continue

        par = doc.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if text.startswith("→"):  # pointer line to a verdict file
            par.paragraph_format.space_before = Pt(2)
            par.paragraph_format.left_indent = Inches(0.14)
        add_runs(par, text)
        last_was_image = False

    doc.save(out)
    print("wrote", out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("source", nargs="?", default="experiments/spectral_graph/INSIGHTS.md")
    ap.add_argument("-o", "--out")
    args = ap.parse_args()

    src = args.source if os.path.isabs(args.source) else os.path.join(REPO, args.source)
    default_name = {
        "INSIGHTS.md": "QML_Insights_Six_Rules.docx",
        "EXPERIMENT_INDEX.md": "QML_Experiment_Index.docx",
    }.get(os.path.basename(src), os.path.splitext(os.path.basename(src))[0] + ".docx")
    out = args.out or os.path.join(os.path.dirname(src), default_name)

    convert(src, out)


if __name__ == "__main__":
    main()
